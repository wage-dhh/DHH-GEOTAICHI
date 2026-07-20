from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re

ROOT = Path(r'F:\dhh\GeoTaichi-dhh')
MD = ROOT / 'docs' / 'kohler_2022_model_alignment_report.md'
OUT = ROOT / 'docs' / 'kohler_2022_model_alignment_report_clean.docx'

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(25, 35, 45)
MUTED = RGBColor(90, 98, 110)
FILL_LIGHT = 'F2F4F7'
BORDER = 'B8C2CC'


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = 'Calibri'
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def style_doc(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ('Heading 1', 16, BLUE, 16, 8),
        ('Heading 2', 13, BLUE, 12, 6),
        ('Heading 3', 12, DARK_BLUE, 8, 4),
    ]:
        st = styles[name]
        st.font.name = 'Calibri'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for name in ['List Bullet', 'List Number']:
        st = styles[name]
        st.font.name = 'Calibri'
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.5)
        st.paragraph_format.first_line_indent = Inches(-0.25)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.10

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run('Kohler et al. (2022) MPM Framework Alignment')
    set_run_font(hr, size=9, color=MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run('GeoTaichi-dhh | Model Alignment Report')
    set_run_font(fr, size=9, color=MUTED)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table, color=BORDER, size='6'):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        element = borders.find(qn(f'w:{edge}'))
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_table_width(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths)))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_ind = tbl_pr.first_child_found_in('w:tblInd')
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), '120')
    tbl_ind.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement('w:tblGrid')
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(w))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn('w:tcW'))
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                tc_pr.append(tc_w)
            tc_w.set(qn('w:w'), str(widths[i]))
            tc_w.set(qn('w:type'), 'dxa')
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_inline_markup(paragraph, text):
    # Handles simple markdown code spans and bold spans.
    parts = re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = DARK_BLUE
        elif part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def add_table_from_md(doc, lines):
    rows = []
    for line in lines:
        if not line.strip().startswith('|'):
            continue
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        if all(set(c) <= set('-: ') for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [max(1200, int(9360 / cols))] * cols
    if cols == 3:
        widths = [2400, 1600, 5360]
    elif cols == 2:
        widths = [3000, 6360]
    elif cols == 4:
        widths = [1800, 1400, 3000, 3160]
    set_table_width(table, widths)
    set_table_borders(table)
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            text = row[c_idx] if c_idx < len(row) else ''
            if r_idx == 0:
                set_cell_shading(cell, FILL_LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_markup(p, text)
            for run in p.runs:
                if r_idx == 0:
                    run.font.bold = True
                    run.font.color.rgb = DARK_BLUE
                run.font.size = Pt(9.5)
    doc.add_paragraph()


def add_code_block(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, [9360])
    set_table_borders(table, color='D0D7DE', size='4')
    cell = table.cell(0, 0)
    set_cell_shading(cell, 'F6F8FA')
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(36, 41, 47)
    doc.add_paragraph()


def build_doc():
    text = MD.read_text(encoding='utf-8')
    doc = Document()
    style_doc(doc)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run('Kohler et al. (2022) MPM Framework Alignment Report')
    set_run_font(r, size=24, bold=True, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run('GeoTaichi-dhh code alignment assessment against the paper model')
    set_run_font(r, size=13, color=MUTED)
    p = doc.add_paragraph()
    add_inline_markup(p, 'Date: 2026-07-16')
    p = doc.add_paragraph()
    add_inline_markup(p, 'Repository: F:\\dhh\\GeoTaichi-dhh')
    p = doc.add_paragraph()
    add_inline_markup(p, 'Assessment: equivalent boundary validation mostly aligned; full strict reproduction remains partial.')
    doc.add_page_break()

    lines = text.splitlines()
    i = 0
    in_code = False
    code_lines = []
    table_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith('```'):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                add_code_block(doc, '\n'.join(code_lines))
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if stripped.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            add_table_from_md(doc, table_lines)
            continue
        if stripped == '':
            i += 1
            continue
        if stripped.startswith('# '):
            # Already have cover title; add as normal H1 if it is not duplicate.
            if 'Kohler et al.' not in stripped:
                doc.add_heading(stripped[2:], level=1)
            i += 1
            continue
        if stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=1)
            i += 1
            continue
        if stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=2)
            i += 1
            continue
        if stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_inline_markup(p, stripped[2:])
            i += 1
            continue
        m = re.match(r'(\d+)\.\s+(.*)', stripped)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_inline_markup(p, m.group(2))
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline_markup(p, stripped)
        i += 1

    doc.save(OUT)
    return OUT

if __name__ == '__main__':
    out = build_doc()
    print(out)
