from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r'F:\dhh\GeoTaichi-dhh\docs\kohler_2022_model_alignment_report_chinese.docx')

def U(s):
    return s.encode('ascii').decode('unicode_escape')

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(25, 35, 45)
MUTED = RGBColor(90, 98, 110)
BORDER = 'B8C2CC'
FILL = 'F2F4F7'
GOOD = 'E7F4EA'
PARTIAL = 'FFF4CC'
BAD = 'FCE8E6'
CALLOUT = 'E8EEF5'


def font(run, size=None, bold=None, color=None):
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color: run.font.color.rgb = color


def style_doc(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in [('Heading 1',16,BLUE,16,8),('Heading 2',13,BLUE,12,6),('Heading 3',12,DARK_BLUE,8,4)]:
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name in ['List Bullet', 'List Number']:
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.5)
        st.paragraph_format.first_line_indent = Inches(-0.25)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.10
    h = sec.header.paragraphs[0]
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = h.add_run(U(r'Kohler 2022 MPM \u5bf9\u9f50\u6027\u62a5\u544a'))
    font(r, 9, False, MUTED)
    f = sec.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = f.add_run(U(r'GeoTaichi-dhh | \u4e2d\u6587\u62a5\u544a'))
    font(r, 9, False, MUTED)


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        pr.append(shd)
    shd.set(qn('w:fill'), fill)


def margins(cell):
    pr = cell._tc.get_or_add_tcPr()
    mar = pr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar')
        pr.append(mar)
    for k,v in [('top',80),('start',120),('bottom',80),('end',120)]:
        node = mar.find(qn('w:'+k))
        if node is None:
            node = OxmlElement('w:'+k)
            mar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')


def borders(table):
    pr = table._tbl.tblPr
    b = pr.first_child_found_in('w:tblBorders')
    if b is None:
        b = OxmlElement('w:tblBorders'); pr.append(b)
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e = b.find(qn('w:'+edge))
        if e is None:
            e = OxmlElement('w:'+edge); b.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:space'),'0'); e.set(qn('w:color'),BORDER)


def table_width(table, widths):
    table.autofit = False
    pr = table._tbl.tblPr
    tw = pr.first_child_found_in('w:tblW')
    if tw is None:
        tw = OxmlElement('w:tblW'); pr.append(tw)
    tw.set(qn('w:w'), str(sum(widths))); tw.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement('w:tblGrid'); table._tbl.insert(0, grid)
    for c in list(grid): grid.remove(c)
    for w in widths:
        col = OxmlElement('w:gridCol'); col.set(qn('w:w'), str(w)); grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            pr = cell._tc.get_or_add_tcPr()
            cw = pr.find(qn('w:tcW'))
            if cw is None:
                cw = OxmlElement('w:tcW'); pr.append(cw)
            cw.set(qn('w:w'), str(widths[i])); cw.set(qn('w:type'), 'dxa')
            margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_p(doc, text, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    font(r, 11, bold, color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(text)
    font(r)


def add_callout(doc, title, body, fill=CALLOUT):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_width(t, [9360]); borders(t)
    c = t.cell(0,0); shade(c, fill)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title); font(r, 11, True, DARK_BLUE)
    p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(0); p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body); font(r2)
    doc.add_paragraph()


def fill_table(table, headers, rows, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_width(table, widths); borders(table)
    for i,h in enumerate(headers):
        c = table.cell(0,i); shade(c,FILL)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h); font(r, 10, True, DARK_BLUE)
    for ri,row in enumerate(rows,1):
        for ci,val in enumerate(row):
            c = table.cell(ri,ci)
            if ci == 1:
                s = str(val)
                if U(r'\u4e0d') in s or 'Partial' in s: shade(c, BAD if U(r'\u4e0d') in s else PARTIAL)
                elif U(r'\u90e8\u5206') in s: shade(c, PARTIAL)
                elif U(r'\u5bf9\u9f50') in s or U(r'\u901a\u8fc7') in s: shade(c, GOOD)
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(val)); font(r, 9.5)


def main():
    doc = Document(); style_doc(doc)
    title = U(r'Kohler et al. (2022) MPM \u6846\u67b6\u5bf9\u9f50\u6027\u8bc4\u4f30\u62a5\u544a')
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(80); p.paragraph_format.space_after = Pt(10)
    r = p.add_run(title); font(r, 24, True, DARK_BLUE)
    p = doc.add_paragraph(); r = p.add_run(U(r'GeoTaichi-dhh \u5f53\u524d\u4ee3\u7801\u4e0e\u8bba\u6587\u4e3b\u8981\u6a21\u578b\u7684\u5bf9\u9f50\u68c0\u67e5'))
    font(r, 13, False, MUTED)
    add_p(doc, U(r'\u65e5\u671f: 2026-07-16'))
    add_p(doc, U(r'\u4ed3\u5e93: F:\\dhh\\GeoTaichi-dhh'))
    add_callout(doc, U(r'\u6838\u5fc3\u7ed3\u8bba'), U(r'\u5f53\u524d\u4ee3\u7801\u4e0e Kohler et al. (2022) \u7684 free-field / quiet boundary \u601d\u60f3\u57fa\u672c\u5bf9\u9f50, \u4f46\u4e0d\u662f\u8bba\u6587\u5b8c\u6574 MPM \u6846\u67b6\u6216 FLAC3D Example 3.3 apply-ff \u7684\u4e25\u683c\u590d\u73b0.'))
    doc.add_page_break()

    doc.add_heading(U(r'1. \u6267\u884c\u6458\u8981'), level=1)
    add_p(doc, U(r'\u672c\u62a5\u544a\u68c0\u67e5 compliant base boundary, lateral free-field boundary, free-field column, \u52a8\u6001\u5e94\u529b\u63d0\u53d6, dashpot \u8026\u5408\u548c Example 3.3 \u51e0\u4f55\u590d\u73b0.'))
    add_p(doc, U(r'\u6700\u5f3a\u7684\u5bf9\u9f50\u70b9\u662f reduced Kohler-style \u8fb9\u754c\u9a8c\u8bc1\u6a21\u578b: \u72ec\u7acb\u5de6\u53f3\u81ea\u7531\u573a\u67f1, \u9759\u6001\u652f\u6491, \u52a8\u6001\u81ea\u7531\u573a\u5e94\u529b, \u76f8\u5bf9\u901f\u5ea6 dashpot \u548c\u89d2\u70b9\u53e0\u52a0\u90fd\u5df2\u5b9e\u73b0.'))
    add_p(doc, U(r'\u4e3b\u8981\u4e0d\u8db3\u662f: \u5f53\u524d\u8def\u5f84\u4ecd\u662f\u7b49\u6548 2D \u526a\u5207\u6ce2\u9a8c\u8bc1, \u4e0d\u662f\u5b8c\u6574 GeoTaichi MPM \u4e3b\u4f53\u6c42\u89e3\u5668\u7684\u901a\u7528\u8fb9\u754c\u6761\u4ef6.'))

    doc.add_heading(U(r'2. \u5bf9\u9f50\u77e9\u9635'), level=1)
    rows = [
        (U(r'\u6750\u6599\u6ce2\u901f\u4e0e dashpot \u7cfb\u6570'), U(r'\u5bf9\u9f50'), U(r'Cs/Cp \u548c eta_s=rho*Cs, eta_p=rho*Cp \u516c\u5f0f\u6b63\u786e.')),
        (U(r'\u72ec\u7acb free-field columns'), U(r'\u57fa\u672c\u5bf9\u9f50'), U(r'\u5de6\u53f3\u81ea\u7531\u573a\u67f1\u5177\u6709\u72ec\u7acb\u8d28\u91cf, \u52a8\u91cf, \u901f\u5ea6, \u5e94\u529b\u548c\u53d8\u5f62\u72b6\u6001.')),
        (U(r'\u5468\u671f\u81ea\u7531\u573a\u67f1'), U(r'\u90e8\u5206\u5bf9\u9f50'), U(r'\u5f53\u524d\u91c7\u7528\u5de6\u53f3\u8282\u70b9\u5e73\u5747, \u4e0d\u662f\u540c\u7f16\u53f7/\u540c\u5185\u5b58\u62d3\u6251.')),
        (U(r'Eq.30 \u4fa7\u5411\u7275\u5f15\u5206\u89e3'), U(r'\u57fa\u672c\u5bf9\u9f50'), U(r'\u663e\u5f0f\u7ec4\u5408 static support, dynamic stress \u548c dashpot relative velocity.')),
        (U(r'Eq.31 \u52a8\u6001\u5e94\u529b'), U(r'\u5bf9\u9f50'), U(r'\u4f7f\u7528 sigma_current - sigma_static.')),
        (U(r'\u5b8c\u6574 MPM \u4e3b\u6a21\u578b'), U(r'\u4e0d\u5bf9\u9f50'), U(r'\u5f53\u524d\u9a8c\u8bc1\u8def\u5f84\u7ed5\u8fc7\u5b8c\u6574 GeoTaichi MPM \u4e3b\u4f53\u6c42\u89e3\u5668.')),
        (U(r'FLAC3D brick/wedge \u51e0\u4f55'), U(r'\u4e0d\u5bf9\u9f50'), U(r'\u5f53\u524d\u662f hand-coded envelope / particle cloud, \u4e0d\u4fdd\u7559 zone connectivity \u548c GP ID.')),
        (U(r'Eq.32 dp \u9762\u79ef\u66f4\u65b0'), U(r'\u4e0d\u5bf9\u9f50'), U(r'\u5f53\u524d\u4f7f\u7528 delta_p=dz, \u672a\u6309 deformation-gradient stretch \u66f4\u65b0.')),
        (U(r'\u672c\u6784\u6a21\u578b'), U(r'\u90e8\u5206\u5bf9\u9f50'), U(r'\u5f53\u524d\u662f\u7ebf\u5f39\u6027\u526a\u5207\u54cd\u5e94, \u4e0d\u662f\u5b8c\u6574\u5f39\u5851\u6027\u5e94\u529b\u79ef\u5206.')),
    ]
    t = doc.add_table(rows=1+len(rows), cols=3)
    fill_table(t, [U(r'\u7ec4\u4ef6'), U(r'\u72b6\u6001'), U(r'\u8bf4\u660e')], rows, [2600, 1500, 5260])

    doc.add_heading(U(r'3. \u5df2\u7ecf\u5bf9\u9f50\u7684\u5185\u5bb9'), level=1)
    for txt in [
        U(r'dashpot \u7cfb\u6570\u4e0e\u8bba\u6587\u4e00\u81f4: eta_s = rho * Cs, eta_p = rho * Cp.'),
        U(r'\u81ea\u7531\u573a\u67f1\u72ec\u7acb\u66f4\u65b0, \u4e3b\u6a21\u578b\u4e0d\u53cd\u5411\u5f71\u54cd\u81ea\u7531\u573a.'),
        U(r'\u4fa7\u5411\u7275\u5f15\u5305\u542b\u9759\u6001\u652f\u6491, \u52a8\u6001\u81ea\u7531\u573a\u5e94\u529b\u548c\u76f8\u5bf9\u901f\u5ea6 dashpot.'),
        U(r'\u89d2\u70b9\u5904 base \u548c lateral contribution \u76f4\u63a5\u53e0\u52a0, \u4e0e\u8bba\u6587\u63cf\u8ff0\u4e00\u81f4.'),
    ]:
        add_bullet(doc, txt)

    doc.add_heading(U(r'4. \u4e3b\u8981\u5dee\u8ddd'), level=1)
    for txt in [
        U(r'\u4e0d\u662f\u4e25\u683c FLAC3D zone-to-zone reproduction; \u53ea\u80fd\u8bf4\u662f\u7b49\u6548\u8fb9\u754c\u9a8c\u8bc1.'),
        U(r'\u672a\u5b9e\u73b0 FLAC3D brick/wedge zone mesh \u548c apply-ff side/corner grid.'),
        U(r'\u672a\u5c06\u8fb9\u754c\u903b\u8f91\u5b8c\u6574\u63a5\u5165\u751f\u4ea7 MPM \u6c42\u89e3\u5668\u5faa\u73af.'),
        U(r'\u5f62\u51fd\u6570\u4e0e\u8bba\u6587\u4e0d\u540c: \u5f53\u524d\u4e3a\u7ebf\u6027/\u53cc\u7ebf\u6027, \u8bba\u6587\u4e3a cubic B-spline.'),
        U(r'\u8fb9\u754c\u9762\u79ef dp \u672a\u6309 Eq.32 \u7528 deformation gradient \u66f4\u65b0.'),
    ]:
        add_bullet(doc, txt)

    doc.add_heading(U(r'5. \u5f53\u524d\u9a8c\u8bc1\u7ed3\u679c'), level=1)
    add_callout(doc, U(r'\u73b0\u6709\u8f93\u51fa\u72b6\u6001'), U(r'Kohler-style MPM boundary validation against FLAC3D Example 3.3 = Successful\nStrict FLAC3D apply-ff reproduction = Partial'), PARTIAL)
    metrics = [
        ('main correlation', '0.9835286353906073'),
        ('main NRMSE', '0.13445657938040273'),
        ('main peak error percent', '-19.18738532462333'),
        ('main phase lag', '0.0'),
    ]
    t = doc.add_table(rows=1+len(metrics), cols=2)
    fill_table(t, [U(r'\u6307\u6807'), U(r'\u5f53\u524d\u503c')], metrics, [3000, 6360])

    doc.add_heading(U(r'6. \u5efa\u8bae'), level=1)
    for txt in [
        U(r'\u4fdd\u7559\u4e24\u4e2a\u72b6\u6001: \u7b49\u6548 Kohler-style boundary validation \u548c strict FLAC3D/Kohler reproduction.'),
        U(r'\u589e\u52a0\u4e25\u683c\u51e0\u4f55\u8def\u5f84: FLAC3D command parser \u6216 VTU unstructured mesh import.'),
        U(r'\u5c06 reduced boundary manager \u63a8\u8fdb\u4e3a\u5b8c\u6574 MPM \u6c42\u89e3\u5668\u5185\u7684\u901a\u7528\u8fb9\u754c\u6761\u4ef6.'),
        U(r'\u6309 Eq.32 \u5b9e\u73b0 deformation-gradient-based boundary area update.'),
        U(r'\u5982\u9700\u4e25\u683c\u5bf9\u9f50\u8bba\u6587, \u9700\u8865\u9f50 cubic B-spline \u548c mirrored-particle boundary handling.'),
    ]:
        add_bullet(doc, txt)

    doc.add_heading(U(r'7. \u6700\u7ec8\u8bc4\u4f30'), level=1)
    final_rows = [
        (U(r'\u8fb9\u754c\u673a\u5236\u5bf9\u9f50'), U(r'\u57fa\u672c\u5bf9\u9f50'), U(r'\u9002\u7528\u4e8e reduced shear-wave validation.')),
        (U(r'\u5b8c\u6574 Kohler 2022 \u6846\u67b6'), U(r'\u90e8\u5206\u5bf9\u9f50'), U(r'\u7f3a\u5c11\u5b8c\u6574\u4e3b\u6a21\u578b, \u5927\u53d8\u5f62 dp \u66f4\u65b0\u548c\u5b8c\u6574\u672c\u6784.')),
        (U(r'FLAC3D Example 3.3 apply-ff'), U(r'Partial / \u4e0d\u4e25\u683c'), U(r'\u7f3a\u5c11 brick/wedge topology \u548c apply-ff grid reproduction.')),
    ]
    t = doc.add_table(rows=1+len(final_rows), cols=3)
    fill_table(t, [U(r'\u8bc4\u4f30\u9879'), U(r'\u7ed3\u8bba'), U(r'\u8bf4\u660e')], final_rows, [2600, 1900, 4860])
    add_callout(doc, U(r'\u4e00\u53e5\u8bdd\u7ed3\u8bba'), U(r'\u5f53\u524d\u4ee3\u7801\u9002\u5408\u5c55\u793a\u548c\u8c03\u8bd5 Kohler-style \u8fb9\u754c\u65b9\u7a0b; \u82e5\u76ee\u6807\u662f\u8bba\u6587\u5b8c\u6574\u6846\u67b6\u6216 FLAC3D apply-ff \u4e25\u683c\u590d\u73b0, \u8fd8\u9700\u8981\u8865\u9f50\u51e0\u4f55\u62d3\u6251, \u5b8c\u6574 MPM \u96c6\u6210, \u5f62\u51fd\u6570, dp \u66f4\u65b0\u548c\u9759\u529b\u5e94\u529b\u8f6c\u79fb.'))
    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    main()
